from django.utils import timezone
from django.http import HttpResponse
from io import BytesIO
from rest_framework import viewsets, permissions
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied
from rest_framework.response import Response
from .models import Invoice
from .serializers import InvoiceSerializer


def infer_role(user):
    group = user.groups.first()
    if group:
        role = group.name.lower()
        if role in ["gp", "nurse", "admin", "superadmin", "patient"]:
            return role
    username = (user.username or "").lower()
    if username == "superadmin":
        return "superadmin"
    if username.startswith("admin"):
        return "admin"
    if username.startswith("nurse"):
        return "nurse"
    if hasattr(user, "patient_profile") and user.patient_profile is not None:
        return "patient"
    return "gp"


def _format_invoice_datetime(value):
    if not value:
        return "-"
    if timezone.is_naive(value):
        return value.strftime("%Y-%m-%d %H:%M")
    return timezone.localtime(value).strftime("%Y-%m-%d %H:%M")


def _invoice_text_payload(invoice):
    issued = _format_invoice_datetime(invoice.issued_at)
    paid = _format_invoice_datetime(invoice.paid_at)
    patient_name = f"{invoice.patient.first_name} {invoice.patient.last_name}".strip()
    return {
        "invoice_id": str(invoice.id),
        "patient_name": patient_name,
        "description": invoice.description,
        "amount": f"{invoice.amount}",
        "status": invoice.status,
        "issued": issued,
        "paid": paid,
    }


def _build_invoice_docx(invoice):
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required for Word export.") from exc

    payload = _invoice_text_payload(invoice)
    doc = Document()
    doc.add_heading("SHMS Invoice Summary", level=0)
    doc.add_paragraph(f"Invoice ID: {payload['invoice_id']}")

    info = doc.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    for key, value in [
        ("Patient", payload["patient_name"]),
        ("Description", payload["description"]),
        ("Amount", f"${payload['amount']}"),
        ("Status", payload["status"].upper()),
        ("Issued", payload["issued"]),
        ("Paid", payload["paid"]),
    ]:
        row = info.add_row().cells
        row[0].text = key
        row[1].text = value

    doc.add_paragraph("Thank you for choosing SHMS Healthcare Services.")
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _build_invoice_pdf(invoice):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
    except ImportError as exc:
        raise RuntimeError("reportlab is required for PDF export.") from exc

    payload = _invoice_text_payload(invoice)
    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4

    c.setFillColor(colors.HexColor("#1e3a8a"))
    c.rect(0, height - 110, width, 110, stroke=0, fill=1)
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 24)
    c.drawString(44, height - 58, "SHMS")
    c.setFont("Helvetica", 12)
    c.drawString(44, height - 78, "Invoice Summary")

    y = height - 150
    c.setFillColor(colors.black)
    c.setFont("Helvetica", 11)
    lines = [
        ("Invoice ID", payload["invoice_id"]),
        ("Patient", payload["patient_name"]),
        ("Description", payload["description"]),
        ("Amount", f"${payload['amount']}"),
        ("Status", payload["status"].upper()),
        ("Issued", payload["issued"]),
        ("Paid", payload["paid"]),
    ]
    for label, value in lines:
        c.setFont("Helvetica-Bold", 11)
        c.drawString(44, y, f"{label}:")
        c.setFont("Helvetica", 11)
        c.drawString(150, y, str(value))
        y -= 24

    c.setFillColor(colors.HexColor("#334155"))
    c.setFont("Helvetica-Oblique", 10)
    c.drawString(44, 36, "Thank you for choosing SHMS Healthcare Services")
    c.showPage()
    c.save()
    buf.seek(0)
    return buf.getvalue()

class InvoiceViewSet(viewsets.ModelViewSet):
    queryset = Invoice.objects.all()
    serializer_class = InvoiceSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        role = infer_role(self.request.user)
        qs = Invoice.objects.select_related("patient").all()
        if role == "patient":
            qs = qs.filter(patient__user=self.request.user)
        return qs

    @action(detail=True, methods=["post"])
    def mark_paid(self, request, pk=None):
        invoice = self.get_object()
        role = infer_role(request.user)
        if role == "patient" and invoice.patient.user_id != request.user.id:
            raise PermissionDenied("You can only pay your own invoices.")
        invoice.status = "paid"
        invoice.paid_at = timezone.now()
        invoice.save()
        return Response({"status": "paid"})

    @action(detail=True, methods=["get"], url_path="download")
    def download(self, request, pk=None):
        invoice = self.get_object()
        role = infer_role(request.user)
        if role == "patient" and invoice.patient.user_id != request.user.id:
            raise PermissionDenied("You can only download your own invoices.")

        export_format = (
            request.query_params.get("file_type")
            or request.query_params.get("format")
            or "txt"
        ).strip().lower()
        filename_base = f"invoice-{invoice.id}"

        if export_format == "docx":
            try:
                content = _build_invoice_docx(invoice)
            except RuntimeError as exc:
                return Response({"detail": str(exc)}, status=500)
            response = HttpResponse(
                content,
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            response["Content-Disposition"] = f'attachment; filename="{filename_base}.docx"'
            return response

        if export_format == "pdf":
            try:
                content = _build_invoice_pdf(invoice)
            except RuntimeError as exc:
                return Response({"detail": str(exc)}, status=500)
            response = HttpResponse(content, content_type="application/pdf")
            response["Content-Disposition"] = f'attachment; filename="{filename_base}.pdf"'
            return response

        payload = _invoice_text_payload(invoice)
        content = (
            "SHMS Invoice Summary\n"
            "=====================\n"
            f"Invoice ID: {payload['invoice_id']}\n"
            f"Patient: {payload['patient_name']}\n"
            f"Description: {payload['description']}\n"
            f"Amount: ${payload['amount']}\n"
            f"Status: {payload['status']}\n"
            f"Issued At: {payload['issued']}\n"
            f"Paid At: {payload['paid']}\n"
        )

        response = HttpResponse(content, content_type="text/plain; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{filename_base}.txt"'
        return response
